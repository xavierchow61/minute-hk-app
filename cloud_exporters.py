"""Markdown → Word + PDF for cloud (no MS Word needed)"""
import io
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ============================================================
# Word (.docx) — python-docx, pure python, works on cloud
# ============================================================

def md_to_docx_bytes(md_text: str) -> bytes:
    """將 Markdown 轉 Word .docx，返回 bytes"""
    doc = Document()

    # 預設字型（繁中要支援嘅）
    style = doc.styles["Normal"]
    style.font.name = "Microsoft JhengHei UI"
    style.font.size = Pt(11)

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    lines = md_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # Headings
        if line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)

        # Table
        elif line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            if len(table_lines) >= 2:
                header_cells = [c.strip() for c in table_lines[0].strip("|").split("|")]
                rows = []
                for tline in table_lines[2:]:
                    cells = [c.strip() for c in tline.strip("|").split("|")]
                    if len(cells) == len(header_cells):
                        rows.append(cells)
                table = doc.add_table(rows=1 + len(rows), cols=len(header_cells))
                table.style = "Light Grid Accent 1"
                for j, h in enumerate(header_cells):
                    cell = table.rows[0].cells[j]
                    cell.text = h
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.bold = True
                for ridx, row_data in enumerate(rows):
                    for cidx, cval in enumerate(row_data):
                        table.rows[ridx + 1].cells[cidx].text = cval
            continue

        # Quote
        elif line.startswith("> "):
            p = doc.add_paragraph(line[2:])
            p.paragraph_format.left_indent = Cm(1)
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        # Bullet
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:], style="List Bullet")

        # Numbered
        elif re.match(r"^\d+\.\s", line):
            content = re.sub(r"^\d+\.\s", "", line)
            doc.add_paragraph(content, style="List Number")

        # Horizontal rule
        elif line.strip() == "---":
            doc.add_paragraph("─" * 50).alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Blank
        elif not line.strip():
            pass

        # Normal paragraph (handle **bold**)
        else:
            p = doc.add_paragraph()
            parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    p.add_run(part[2:-2]).bold = True
                elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                    p.add_run(part[1:-1]).italic = True
                else:
                    p.add_run(part)
        i += 1

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ============================================================
# PDF — WeasyPrint (Chinese support, requires packages.txt)
# Falls back to a simple text-only PDF if WeasyPrint fails
# ============================================================

def md_to_pdf_bytes(md_text: str) -> bytes:
    """將 Markdown 轉 PDF。優先用 WeasyPrint（要 packages.txt），fallback 用 fpdf2。"""
    try:
        return _md_to_pdf_weasy(md_text)
    except Exception as e:
        # Fallback: 簡單 PDF（純文字 + 基本格式）
        print(f"⚠️ WeasyPrint failed ({e}), using fallback")
        return _md_to_pdf_simple(md_text)


def _md_to_pdf_weasy(md_text: str) -> bytes:
    """WeasyPrint - 靚輸出 + Chinese 支援"""
    import markdown
    from weasyprint import HTML, CSS

    html_body = markdown.markdown(
        md_text,
        extensions=["tables", "fenced_code", "nl2br"],
    )

    html = f"""<!DOCTYPE html>
<html lang="zh-Hant">
<head>
<meta charset="utf-8">
<title>會議紀要</title>
</head>
<body>
{html_body}
</body>
</html>"""

    css = CSS(string="""
        @page {
            size: A4;
            margin: 2cm 2.5cm;
            @bottom-center {
                content: "Page " counter(page) " of " counter(pages);
                font-size: 9pt;
                color: #999;
            }
        }
        body {
            font-family: 'Noto Sans CJK TC', 'Noto Sans CJK SC', 'Microsoft JhengHei',
                         'PingFang TC', 'Heiti TC', sans-serif;
            font-size: 11pt;
            line-height: 1.7;
            color: #1e293b;
        }
        h1 {
            color: #1e66f5;
            font-size: 20pt;
            border-bottom: 2px solid #1e66f5;
            padding-bottom: 8pt;
            margin-bottom: 16pt;
        }
        h2 {
            color: #1e293b;
            font-size: 14pt;
            margin-top: 18pt;
            margin-bottom: 8pt;
            padding-bottom: 4pt;
            border-bottom: 1px solid #e2e8f0;
        }
        h3 {
            color: #475569;
            font-size: 12pt;
            margin-top: 14pt;
            margin-bottom: 6pt;
        }
        p { margin: 6pt 0; }

        /* Table 美化 */
        table {
            width: 100%;
            border-collapse: collapse;
            margin: 12pt 0;
            font-size: 10pt;
        }
        th, td {
            border: 1px solid #cbd5e1;
            padding: 6pt 10pt;
            text-align: left;
        }
        th {
            background: #f1f5f9;
            font-weight: 700;
            color: #1e293b;
        }
        tr:nth-child(even) td { background: #fafbfc; }

        /* 🎯 List：用 custom bullet (唔依賴字體) */
        ul {
            list-style: none;
            padding-left: 0;
            margin: 8pt 0;
        }
        ul li {
            position: relative;
            padding-left: 18pt;
            margin: 5pt 0;
        }
        ul li::before {
            content: "▸";
            color: #1e66f5;
            font-weight: bold;
            position: absolute;
            left: 2pt;
            top: 0;
            font-family: 'Helvetica', 'Arial', sans-serif;
        }
        ol {
            padding-left: 22pt;
            margin: 8pt 0;
        }
        ol li { margin: 5pt 0; padding-left: 4pt; }
        ol li::marker { color: #1e66f5; font-weight: 700; }

        /* Nested lists */
        ul ul, ol ol, ul ol, ol ul {
            margin: 4pt 0 4pt 0;
        }
        ul ul li::before { content: "·"; }

        blockquote {
            border-left: 3px solid #1e66f5;
            padding: 4pt 12pt;
            margin: 10pt 0;
            color: #475569;
            font-style: italic;
            background: #f8fafc;
        }
        hr {
            border: none;
            border-top: 1px solid #e2e8f0;
            margin: 14pt 0;
        }
        code {
            background: #f1f5f9;
            padding: 1pt 5pt;
            border-radius: 3pt;
            font-family: 'Consolas', 'Courier New', monospace;
            font-size: 9.5pt;
            color: #1e66f5;
        }
        strong { color: #1e293b; font-weight: 700; }
        em { color: #475569; }
    """)

    out = io.BytesIO()
    HTML(string=html).write_pdf(out, stylesheets=[css])
    return out.getvalue()


def _md_to_pdf_simple(md_text: str) -> bytes:
    """Fallback: 用 fpdf2，純 python，但 Chinese 字體 limited"""
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=20)

    # Try to use a CJK-capable font if available
    try:
        pdf.add_font("NotoCJK", "", "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", uni=True)
        pdf.set_font("NotoCJK", "", 11)
    except Exception:
        pdf.set_font("Helvetica", "", 11)

    for line in md_text.split("\n"):
        line = line.strip()
        if line.startswith("# "):
            pdf.set_font_size(16)
            pdf.cell(0, 10, line[2:], ln=True)
            pdf.set_font_size(11)
        elif line.startswith("## "):
            pdf.set_font_size(13)
            pdf.cell(0, 8, line[3:], ln=True)
            pdf.set_font_size(11)
        elif line.startswith("- ") or line.startswith("* "):
            pdf.cell(0, 7, "  • " + line[2:], ln=True)
        else:
            if line:
                # Remove markdown chars
                clean = re.sub(r"[*`#|]", "", line)
                pdf.multi_cell(0, 7, clean)

    out = io.BytesIO()
    pdf.output(out)
    return out.getvalue()
