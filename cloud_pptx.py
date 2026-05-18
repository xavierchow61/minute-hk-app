"""Generate PowerPoint (.pptx) from meeting summary"""
import io
import json
import re
from datetime import date

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


# ============================================================
# AI: 將 Markdown summary 結構化做 PPT slides
# ============================================================
PPT_STRUCTURE_PROMPT = """從以下會議紀要 extract 結構化資料，轉成 JSON 俾我生成 PPT slides。

返回格式（嚴格 JSON，唔好其他文字）：
{{
  "title": "會議標題（短）",
  "subtitle": "客戶 / 項目 / 日期 等",
  "agenda": ["議題 1", "議題 2", "議題 3"],
  "topics": [
    {{
      "title": "議題名稱",
      "bullets": ["要點 1", "要點 2", "要點 3"]
    }}
  ],
  "decisions": ["決議 1", "決議 2"],
  "action_items": [
    {{"task": "...", "assignee": "...", "deadline": "..."}}
  ],
  "risks": ["風險 1", "風險 2"],
  "next_steps": "簡短一句結論"
}}

要求：
- 每個 bullet 唔好太長（每個 < 60 字）
- agenda 最多 5 個
- topics 最多 5 個，每個 bullets 最多 5 個
- 中文 + 英文都得，保留原文

紀要：
{summary}

只輸出 JSON，無 markdown code fence。"""


def _structure_with_ai(summary_md: str) -> dict:
    """Use Gemini to convert summary to PPT structure"""
    try:
        from ai import call_with_retry, _client, GEMINI_MODEL
        from google.genai import types as gtypes

        client = _client()

        def _gen(model=GEMINI_MODEL):
            return client.models.generate_content(
                model=model,
                contents=PPT_STRUCTURE_PROMPT.format(summary=summary_md),
                config=gtypes.GenerateContentConfig(
                    temperature=0.2,
                    max_output_tokens=4096,
                ),
            )

        response = call_with_retry(_gen, model=GEMINI_MODEL)
        raw = response.text.strip()
        raw = raw.removeprefix("```json").removeprefix("```").removesuffix("```").strip()
        return json.loads(raw)
    except Exception as e:
        # Fallback: simple parsing
        return _parse_markdown_fallback(summary_md)


def _parse_markdown_fallback(md: str) -> dict:
    """Simple parser if AI fails"""
    return {
        "title": "會議紀要",
        "subtitle": date.today().strftime("%Y-%m-%d"),
        "agenda": [],
        "topics": [{
            "title": "會議內容",
            "bullets": [line[2:] for line in md.split("\n") if line.startswith("- ")][:8],
        }],
        "decisions": [],
        "action_items": [],
        "risks": [],
        "next_steps": "請覆核 AI 紀要",
    }


# ============================================================
# 🎨 PPT Generation (python-pptx)
# ============================================================
# Brand colors
COLOR_PRIMARY = RGBColor(0x06, 0xB6, 0xD4)  # cyan
COLOR_DARK = RGBColor(0x1E, 0x29, 0x3B)
COLOR_MUTED = RGBColor(0x71, 0x71, 0x7A)
COLOR_LIGHT_BG = RGBColor(0xFA, 0xFA, 0xFA)


def _add_title_slide(prs: Presentation, title: str, subtitle: str = ""):
    layout = prs.slide_layouts[0]  # Title slide
    slide = prs.slides.add_slide(layout)

    title_box = slide.shapes.title
    if title_box:
        title_box.text = title
        for para in title_box.text_frame.paragraphs:
            for run in para.runs:
                run.font.color.rgb = COLOR_DARK
                run.font.size = Pt(40)
                run.font.bold = True

    if slide.placeholders and len(slide.placeholders) > 1:
        sub = slide.placeholders[1]
        sub.text = subtitle
        for para in sub.text_frame.paragraphs:
            for run in para.runs:
                run.font.color.rgb = COLOR_MUTED
                run.font.size = Pt(18)


def _add_bullet_slide(prs: Presentation, title: str, bullets: list, accent_emoji: str = ""):
    layout = prs.slide_layouts[1]  # Title and content
    slide = prs.slides.add_slide(layout)

    if slide.shapes.title:
        slide.shapes.title.text = f"{accent_emoji} {title}".strip()
        for para in slide.shapes.title.text_frame.paragraphs:
            for run in para.runs:
                run.font.color.rgb = COLOR_DARK
                run.font.size = Pt(28)
                run.font.bold = True

    if len(slide.placeholders) > 1:
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()
        for i, b in enumerate(bullets):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = str(b)
            p.level = 0
            for run in p.runs:
                run.font.size = Pt(16)
                run.font.color.rgb = COLOR_DARK


def _add_action_items_slide(prs: Presentation, items: list):
    layout = prs.slide_layouts[5]  # Title only
    slide = prs.slides.add_slide(layout)

    if slide.shapes.title:
        slide.shapes.title.text = "📋 Action Items"
        for para in slide.shapes.title.text_frame.paragraphs:
            for run in para.runs:
                run.font.color.rgb = COLOR_DARK
                run.font.size = Pt(28)
                run.font.bold = True

    if not items:
        return

    # Table
    rows = len(items) + 1
    cols = 3
    left = Inches(0.5)
    top = Inches(1.5)
    width = Inches(9)
    height = Inches(0.4 + 0.4 * rows)

    table = slide.shapes.add_table(rows, cols, left, top, width, height).table

    # Headers
    headers = ["待辦事項", "負責人", "Deadline"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY
        for para in cell.text_frame.paragraphs:
            for run in para.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(14)
                run.font.bold = True

    # Rows
    for r, item in enumerate(items, start=1):
        table.cell(r, 0).text = str(item.get("task") or "")
        table.cell(r, 1).text = str(item.get("assignee") or "—")
        table.cell(r, 2).text = str(item.get("deadline") or "—")
        for c in range(3):
            for para in table.cell(r, c).text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(12)
                    run.font.color.rgb = COLOR_DARK


def _add_closing_slide(prs: Presentation, next_steps: str):
    layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(layout)

    if slide.shapes.title:
        slide.shapes.title.text = "✅ 下一步"
        for para in slide.shapes.title.text_frame.paragraphs:
            for run in para.runs:
                run.font.color.rgb = COLOR_DARK
                run.font.size = Pt(32)
                run.font.bold = True

    # Add big text in middle
    txt_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(2))
    tf = txt_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = next_steps or "感謝參與會議"
    for run in p.runs:
        run.font.size = Pt(22)
        run.font.color.rgb = COLOR_MUTED


# ============================================================
# Main entry
# ============================================================
def summary_to_pptx_bytes(summary_md: str) -> bytes:
    """將會議紀要轉做 PPT (.pptx) bytes"""
    data = _structure_with_ai(summary_md)

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    # 1. Title slide
    _add_title_slide(prs, data.get("title", "會議紀要"),
                     data.get("subtitle", date.today().isoformat()))

    # 2. Agenda
    agenda = data.get("agenda", [])
    if agenda:
        _add_bullet_slide(prs, "議程", agenda, "📋")

    # 3. Topics (one per topic)
    for topic in data.get("topics", []):
        title = topic.get("title", "")
        bullets = topic.get("bullets", [])
        if title and bullets:
            _add_bullet_slide(prs, title, bullets, "💡")

    # 4. Decisions
    decisions = data.get("decisions", [])
    if decisions:
        _add_bullet_slide(prs, "決議事項", decisions, "✅")

    # 5. Action Items (table)
    action_items = data.get("action_items", [])
    if action_items:
        _add_action_items_slide(prs, action_items)

    # 6. Risks
    risks = data.get("risks", [])
    if risks:
        _add_bullet_slide(prs, "風險與跟進", risks, "⚠️")

    # 7. Closing
    _add_closing_slide(prs, data.get("next_steps", "感謝出席"))

    # Save to bytes
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
